"""Create the input templates (Excel workbook with Tables + header-only CSVs).

Usage (from dashboard/):  python scripts/make_templates.py
Output: data/templates/energy_input_TEMPLATE.xlsx and data/templates/*_template.csv
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "data" / "templates"
OUT.mkdir(parents=True, exist_ok=True)

SEU_TYPES = ["pump", "compressor", "extruder", "steam_dryer", "cooling", "compressed_air", "boiler", "chiller", "hvac",
             "process_heating", "generic"]
EVENT_TYPES = ["production_turndown", "grade_trial", "equipment_change", "meter_outage", "maintenance", "turnaround",
               "setpoint_change", "process_change"]

SHEETS = {
    "energy_monthly": dict(
        cols=["period", "plant", "seu", "carrier", "energy_gj", "production_ton", "ambient_temp_c", "data_quality"],
        notes=["First day of the month, e.g. 2025-10-01", "Plant code, no spaces (PLANT_A)", "SEU code, no spaces; same code in seu_register",
               "electricity | steam | fuel_gas (must match the tariff keys in the notebook CONFIG)",
               "Energy in GJ. Alternative: replace with energy_value + unit (kWh, MWh, MMBtu, Gcal, GJ, t_steam)",
               "Plant (or line) production in tonnes for that month", "Monthly mean ambient temperature, deg C (leave empty if not relevant)",
               "measured | estimated | missing (leave the energy cell empty when missing)"]),
    "seu_register": dict(
        cols=["plant", "plant_name", "seu", "seu_name", "seu_type", "carrier", "metering", "notes"],
        notes=["Same code as in energy_monthly", "Readable plant name", "Same code as in energy_monthly", "Readable SEU name",
               "Selects the knowledge-base advice: " + ", ".join(SEU_TYPES), "electricity | steam | fuel_gas", "sub-metered | calculated | estimated", "Free text"]),
    "events": dict(
        cols=["plant", "seu", "start", "end", "event_type", "description"],
        notes=["Plant code", "SEU code, or * for the whole plant", "First day of the event", "Last day of the event",
               "e.g. " + ", ".join(EVENT_TYPES) + " (production_turndown gets special treatment)", "What happened, in one sentence"]),
}

HDR = PatternFill("solid", fgColor="1F3A5F")

wb = Workbook()
ws = wb.active
ws.title = "README"
lines = [
    ("Energy performance dashboard: input template", True),
    ("1. Fill the three Tables: energy_monthly (one row per plant x SEU x month), seu_register (one row per SEU), events (optional).", False),
    ("2. Keep the header names exactly. Do not leave example rows in the tables.", False),
    ("3. Use 24+ months for the baseline, plus the months you want to report on (ISO 50006 suggests a representative operating range).", False),
    ("4. Save the workbook into the data/ folder and set energy_file/register_file/events_file in the notebook CONFIG to this file name.", False),
    ("5. Hover over a header cell to read what the column means. Header is in row 1; data starts in row 2.", False),
    ("6. Anonymise or scale values before sending any file to a cloud tool (CONFIG: anonymize, scale_energy, scale_production).", False),
]
for r, (txt, bold) in enumerate(lines, 1):
    ws.cell(r, 1, txt).font = Font(bold=bold, size=14 if bold else 11)
ws.column_dimensions["A"].width = 140

for name, spec in SHEETS.items():
    s = wb.create_sheet(name)
    cols = spec["cols"]
    for j, (c, note) in enumerate(zip(cols, spec["notes"]), 1):
        h = s.cell(1, j, c)
        h.fill, h.font = HDR, Font(bold=True, color="FFFFFF")
        h.comment = Comment(note, "template", width=320, height=90)
        s.column_dimensions[get_column_letter(j)].width = 24
    ref = f"A1:{get_column_letter(len(cols))}2"          # header + one empty row (Excel Tables need a data row)
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    s.add_table(t)
    s.freeze_panes = "A2"

def validation(sheet, header, options):
    idx = SHEETS[sheet]["cols"].index(header) + 1
    col = get_column_letter(idx)
    dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
    wb[sheet].add_data_validation(dv)
    dv.add(f"{col}2:{col}2000")

validation("energy_monthly", "carrier", ["electricity", "steam", "fuel_gas"])
validation("energy_monthly", "data_quality", ["measured", "estimated", "missing"])
validation("seu_register", "carrier", ["electricity", "steam", "fuel_gas"])
validation("seu_register", "seu_type", SEU_TYPES)
validation("events", "event_type", EVENT_TYPES)

u = wb.create_sheet("units")
u.append(["unit", "GJ per unit", "note"])
for row in [("GJ", 1, ""), ("MJ", 0.001, ""), ("kWh", 0.0036, ""), ("MWh", 3.6, ""), ("MMBtu", 1.055056, ""), ("Gcal", 4.1868, ""),
            ("t_steam", "set your own", "depends on steam pressure/temperature and condensate return; set UNIT_TO_GJ in the notebook")]:
    u.append(list(row))
u.column_dimensions["A"].width = 14; u.column_dimensions["B"].width = 16; u.column_dimensions["C"].width = 90

wb.save(OUT / "energy_input_TEMPLATE.xlsx")
for name, spec in SHEETS.items():
    pd.DataFrame(columns=spec["cols"]).to_csv(OUT / f"{name}_template.csv", index=False)
print("wrote", sorted(p.name for p in OUT.iterdir()))
